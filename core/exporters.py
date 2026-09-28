from io import BytesIO
from datetime import datetime

import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)


PRIMARY_ORANGE = "FF9800"
PRIMARY_RED = "FF2B2B"
PRIMARY_TURQUOISE = "18D6B2"
PRIMARY_NAVY = "0B2D4D"
LIGHT_BG = "FFF8F2"
WHITE = "FFFFFF"
GRAY = "E9ECEF"
DARK_TEXT = "1F2937"


def _safe_filename(text):
    text = str(text).strip().lower()

    replacements = {
        " ": "_",
        "/": "_",
        "\\": "_",
        ":": "_",
        ";": "_",
        ",": "_",
        ".": "_",
        "á": "a",
        "é": "e",
        "í": "i",
        "ó": "o",
        "ú": "u",
        "ñ": "n",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text


def _prepare_dataframe(df):
    if df is None:
        return pd.DataFrame()

    df_export = df.copy()

    for col in df_export.columns:
        df_export[col] = df_export[col].apply(
            lambda value: "" if pd.isna(value) else value
        )

    return df_export


def excel_bytes(df, title="PatitaSOFT"):
    df = _prepare_dataframe(df)

    output = BytesIO()

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    total_columns = max(len(df.columns), 1)

    # Título principal
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=total_columns,
    )

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(
        bold=True,
        size=18,
        color=WHITE,
    )
    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=PRIMARY_ORANGE,
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 30

    # Fecha de generación
    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=total_columns,
    )

    date_cell = ws.cell(row=2, column=1)
    date_cell.value = (
        f"Generado por PatitaSOFT - "
        f"{datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    date_cell.font = Font(
        italic=True,
        size=10,
        color=DARK_TEXT,
    )
    date_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[2].height = 22

    # Línea vacía
    header_row = 4

    # Encabezados
    for col_idx, column_name in enumerate(df.columns, start=1):
        cell = ws.cell(
            row=header_row,
            column=col_idx,
            value=str(column_name),
        )

        cell.font = Font(
            bold=True,
            color=WHITE,
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=PRIMARY_NAVY,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Datos
    data_start_row = header_row + 1

    for row_idx, row in enumerate(
        df.itertuples(index=False),
        start=data_start_row,
    ):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.border = thin_border

            if row_idx % 2 == 0:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFF6EE",
                )

    # Bordes para encabezados
    for col_idx in range(1, total_columns + 1):
        ws.cell(
            row=header_row,
            column=col_idx,
        ).border = thin_border

    # Ajustar anchos de columnas
    #
    # Importante:
    # NO usamos col[0].column_letter porque la primera fila
    # contiene celdas combinadas y openpyxl devuelve MergedCell.
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)

        values = []

        # Consideramos encabezado + primeras filas
        max_row_to_check = min(ws.max_row, 204)

        for row_idx in range(header_row, max_row_to_check + 1):
            value = ws.cell(
                row=row_idx,
                column=col_idx,
            ).value

            if value is not None:
                values.append(str(value))

        max_length = max(
            (len(value) for value in values),
            default=0,
        )

        width = min(
            max(max_length + 2, 12),
            42,
        )

        ws.column_dimensions[letter].width = width

    # Congelar encabezados
    ws.freeze_panes = f"A{data_start_row}"

    # Auto filtro
    if len(df.columns) > 0:
        last_letter = get_column_letter(len(df.columns))

        ws.auto_filter.ref = (
            f"A{header_row}:{last_letter}{ws.max_row}"
        )

    # Configuración de impresión
    ws.sheet_view.showGridLines = False

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(output)

    output.seek(0)

    return output.getvalue()


def pdf_bytes(df, title="PatitaSOFT"):
    df = _prepare_dataframe(df)

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=22,
        leftMargin=22,
        topMargin=28,
        bottomMargin=28,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "PatitaSOFTTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0B2D4D"),
        spaceAfter=5,
    )

    subtitle_style = ParagraphStyle(
        "PatitaSOFTSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#6B7280"),
        spaceAfter=12,
    )

    cell_style = ParagraphStyle(
        "PatitaSOFTCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=6.7,
        leading=8,
    )

    header_style = ParagraphStyle(
        "PatitaSOFTHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7,
        leading=8,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    elements = []

    elements.append(
        Paragraph(
            str(title),
            title_style,
        )
    )

    elements.append(
        Paragraph(
            (
                "PatitaSOFT · "
                f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            ),
            subtitle_style,
        )
    )

    elements.append(Spacer(1, 8))

    if df.empty:
        elements.append(
            Paragraph(
                "No hay registros para exportar.",
                styles["Normal"],
            )
        )

        document.build(elements)
        output.seek(0)
        return output.getvalue()

    table_data = []

    # Encabezados
    table_data.append(
        [
            Paragraph(
                str(column),
                header_style,
            )
            for column in df.columns
        ]
    )

    # Filas
    for row in df.itertuples(index=False):
        table_data.append(
            [
                Paragraph(
                    str(value),
                    cell_style,
                )
                for value in row
            ]
        )

    page_width = landscape(A4)[0] - 44

    number_columns = max(len(df.columns), 1)

    default_column_width = page_width / number_columns

    column_widths = [
        default_column_width
        for _ in range(number_columns)
    ]

    table = Table(
        table_data,
        colWidths=column_widths,
        repeatRows=1,
    )

    table_style = TableStyle(
        [
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#0B2D4D"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, 0),
                "CENTER",
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D1D5DB"),
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                4,
            ),
        ]
    )

    # Filas alternadas
    for row_idx in range(1, len(table_data)):
        if row_idx % 2 == 0:
            table_style.add(
                "BACKGROUND",
                (0, row_idx),
                (-1, row_idx),
                colors.HexColor("#FFF8F2"),
            )

    table.setStyle(table_style)

    elements.append(table)

    document.build(elements)

    output.seek(0)

    return output.getvalue()


def export_buttons(df, name="Exportación"):
    df = _prepare_dataframe(df)

    excel_file = excel_bytes(
        df,
        title=name,
    )

    pdf_file = pdf_bytes(
        df,
        title=name,
    )

    filename = _safe_filename(name)

    col_excel, col_pdf = st.columns(2)

    with col_excel:
        st.download_button(
            label="📊 Descargar Excel",
            data=excel_file,
            file_name=f"{filename}.xlsx",
            mime=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )

    with col_pdf:
        st.download_button(
            label="📄 Descargar PDF",
            data=pdf_file,
            file_name=f"{filename}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )